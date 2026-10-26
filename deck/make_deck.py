#!/usr/bin/env python3
# =============================================================================
# deck/make_deck.py -- 10-minute presentation deck (EC4308, Week 13)
# =============================================================================
# Builds deck/EC4308_Presentation.pptx from the computed results:
# figures (output/figures/figure*.png) and tables (output/tables/table*_h1.csv).
# Speaker notes are embedded per slide; each of the four members speaks.
# Run AFTER R/11_evaluate.R and R/13_figures.R:
#     python3 deck/make_deck.py
# =============================================================================
import csv, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "output", "figures")
TAB = os.path.join(ROOT, "output", "tables")

DARK = RGBColor(0x1F, 0x36, 0x5C)
ACC  = RGBColor(0xC0, 0x50, 0x2D)

def load_table(name):
    with open(os.path.join(TAB, name), newline="") as f:
        return list(csv.DictReader(f))

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def add_title(slide, text, sub=None):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(1.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = text
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = DARK
    if sub:
        p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = sub
        r2.font.size = Pt(14); r2.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

def add_bullets(slide, items, size=16, left=0.7, top=1.4, width=11.9, height=5.4):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (txt, lvl) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt; p.level = lvl
        for r in p.runs: r.font.size = Pt(size)

def add_pic(slide, path, left, top, width=None, height=None):
    slide.shapes.add_picture(path, Inches(left), Inches(top),
                             Inches(width) if width else None,
                             Inches(height) if height else None)

def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

# ---- key numbers -------------------------------------------------------------
t2 = load_table("table2_h1.csv")
t3 = load_table("table3_h1.csv")
t5 = load_table("table5_neely_replication.csv")
gr = next(r for r in t2 if r["model"] == "gr_comb")
rl = next(r for r in t2 if r["model"] == "rlasso")
bh = next(r for r in t3 if r["model"] == "Buy-and-Hold")
pm = next(r for r in t3 if r["model"] == "Prevailing Mean")
gr3 = next(r for r in t3 if r["model"] == "gr_comb")
post = next(r for r in t3 if r["model"] == "postlasso")
vol212 = next(r for r in t5 if r["sample"] == "Neely sample 1966-2011" and r["signal"] == "VOL_2_12")
vol312 = next(r for r in t5 if r["sample"] == "Neely sample 1966-2011" and r["signal"] == "VOL_3_12")
ext = next(r for r in t5 if r["sample"] == "Extension 2012-2025" and r["signal"] == "VOL_3_12")

# ---- 1. title -----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
tb = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(2.4))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run()
r.text = "Forecasting the Equity Risk Premium Out-of-Sample:\nRegularized ML, Technical Indicators & Tactical Asset Allocation"
r.font.size = Pt(32); r.font.bold = True; r.font.color.rgb = DARK
p2 = tf.add_paragraph(); r2 = p2.add_run()
r2.text = "EC4308 Machine Learning and Economic Forecasting — Group Project (Week 13)"
r2.font.size = Pt(18); r2.font.color.rgb = ACC
p3 = tf.add_paragraph(); r3 = p3.add_run()
r3.text = "Leon • En Shao • Sasinthiran Kannan • Vienna"
r3.font.size = Pt(18)
add_notes(s, "Vienna (0:45): title, the Goyal-Welch puzzle in one sentence, roadmap of the talk.")

# ---- 2. motivation ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
add_title(s, "Motivation: Can anything beat the historical average?", "Goyal–Welch (2008, 2024) vs Neely, Rapach, Tu & Zhou (2014)")
add_bullets(s, [
    ("The equity premium is the single most important input for portfolio choice — 1 month of lead time is economically valuable.", 0),
    ("Goyal & Welch: essentially no macro fundamental beats the prevailing mean out-of-sample; Goyal–Welch–Zafirov (2024) reach the same conclusion with 57 predictors.", 0),
    ("Neely et al. (2014): 14 technical indicators (MA, MOM, OBV) DO predict the premium out-of-sample on 1966–2011.", 0),
    ("Our question: with modern regularized ML, do technical indicators + macro lags beat the prevailing mean on the extended 1951–2025 sample?", 0),
    ("We replicate Neely et al. exactly (signals match 99.98%), then extend the sample to 2025.", 0),
])
add_notes(s, "Vienna (1:00): set up the puzzle. Emphasise: replication first, extension second.")

# ---- 3. data ------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
add_title(s, "Data: 899 months, 98 predictors, zero lookahead")
add_bullets(s, [
    ("Target: monthly log equity premium (%), y = 100·(ln(1+Rᵐ) − ln(1+Rᶠ)), S&P 500 over the risk-free rate.", 0),
    ("Macro: the 14 Goyal–Welch predictors, each with lags 1–6  →  84 columns.", 0),
    ("Technical: 14 binary Neely signals — 6 MA, 2 MOM, 6 OBV — computed from raw index level + S&P 500 volume.", 0),
    ("Sample: 1951-02 → 2025-12; predictors aligned to 'observable at end of month t' — no lookahead leakage.", 0),
    ("Validation: our computed signals match the published Neely series in 99.98% of signal-months (12/14 exact).", 0),
    ("Fixed 240-month rolling window; first OOS forecast 1971-02; n = 659 OOS months at h = 1.", 0),
])
add_notes(s, "Leon (1:15): data sources, alignment convention, the 99.98% validation number, sample dates.")

# ---- 4. methods ----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
add_title(s, "Methods: 7 models + 2 combinations vs the prevailing mean")
add_bullets(s, [
    ("Benchmark: prevailing (historical) mean — the only admissible benchmark.", 0),
    ("Models (same 98-predictor matrix, re-estimated every window):", 0),
    ("LASSO (hdm, theoretical penalty — no CV) & post-LASSO", 1),
    ("Ridge and Elastic Net (glmnet, 10-fold CV inside the window)", 1),
    ("PCR & PLS (CV-selected components) • Random Forest (500 trees, fixed seed)", 1),
    ("Combinations: equal-weight 1/N and constrained Granger–Ramanathan (Σw=1, w≥0, weights from past OOS errors only).", 0),
    ("Horizons h ∈ {1, 3, 12}; DM (1995) tests with HAC lag h−1; Clark–West as appendix check.", 0),
])
add_notes(s, "En Shao (1:30): one line per method, stress the no-lookahead CV protocol and the GR constraint.")

# ---- 5. replication -------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
add_title(s, "Step 1: we replicate Neely et al. (2014) — their result survives",
          "Expanding-window bivariate regressions, one signal at a time, Clark–West test")
add_bullets(s, [
    (f"On their 1966–2011 sample: VOL_2_12 R²_OOS = +{vol212['R2_OOS_pct']}% (CW p = {vol212['CW_p']}), "
     f"VOL_3_12 = +{vol312['R2_OOS_pct']}% (p = {vol312['CW_p']}) — positive and significant, as published (+0.33%).", 0),
    ("Volume-based (OBV) signals carry the predictive content; MA/MOM are weaker — the same pattern as the paper.", 0),
    ("This validates our data pipeline and test implementation before we extend the sample.", 0),
    (f"Extension 2012–2025: every signal turns negative/insignificant (VOL_3_12 = {ext['R2_OOS_pct']}%, p = {ext['CW_p']}).", 0),
])
add_notes(s, "Sasi (1:15): replication = credibility. Show the numbers side by side with the published ones.")

# ---- 6. statistical results ------------------------------------------------------
s = prs.slides.add_slide(BLANK)
add_title(s, "Result 1: nothing beats the prevailing mean on the extended sample")
add_pic(s, os.path.join(FIG, "figure2_r2oos.png"), 0.6, 1.35, width=6.4)
add_bullets(s, [
    ("All models show NEGATIVE R²_OOS at every horizon: the Goyal–Welch result, reproduced with modern ML.", 0),
    (f"Best at h=1: GR combination R²_OOS = {gr['R2_OOS_pct']}%, rlasso = {rl['R2_OOS_pct']}% (PM RMSE 4.40 vs {gr['RMSE'][:5]}).", 0),
    ("Hit rates 54–58% vs PM 61% — the direction of the premium is also hard to call.", 0),
    ("Post-2000 (and post-2011) predictability decay drives the result: the information is real but transient.", 0),
], left=7.3, top=1.35, width=5.6, height=5.0)
add_notes(s, "Sasi (1:15): negative results are results — correct execution matters more than a positive sign. Walk through the bar chart.")

# ---- 7. economic results -----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
add_title(s, "Result 2: forecasts still have economic value — drawdown control")
add_pic(s, os.path.join(FIG, "figure1_wealth.png"), 0.6, 1.4, width=7.2)
add_bullets(s, [
    (f"Switching (10 bps): model strategies cap the max drawdown at {post['sw_maxdd_10bps'][1:]} vs buy-and-hold −0.55; terminal wealth ${post['sw_wealth_10bps']} vs ${bh['sw_wealth_bh']} for buy-and-hold.", 0),
    (f"Mean–variance investor (γ=3): post-LASSO ΔU = +{post['mv_deltaU_10bps_pct']}% p.a., rlasso +{next(r['mv_deltaU_10bps_pct'] for r in t3 if r['model']=='rlasso')}% p.a., GR +{gr3['mv_deltaU_10bps_pct']}% p.a. (10 bps net).", 0),
    ("Forecasts earn their keep as risk-management signals, not as point-forecast improvements.", 0),
], left=8.0, top=1.4, width=4.9, height=5.0)
add_notes(s, "Sasi (1:15): Sharpe vs wealth trade-off; ΔU interpretation; fee sensitivity 5/20 bps in the report.")

# ---- 8. interpretation -------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
add_title(s, "Why did predictability decay?")
add_bullets(s, [
    ("Recession months drive the signal value: volume/MA rules earn their R² in a handful of crisis months (2000–02, 2008–09).", 0),
    ("Post-2011: fewer deep drawdowns, a passive-dominated market, and faster arbitrage of published signals.", 0),
    ("Regularized multivariate fits overfit relative to the simple bivariate design — shrinkage toward the mean is aggressive.", 0),
    ("Consistent with Goyal–Welch–Zafirov (2024): the premium remains extremely hard to predict out-of-sample.", 0),
])
add_notes(s, "Vienna (1:00): economic interpretation, connect to the literature.")

# ---- 9. conclusion -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
add_title(s, "Conclusion")
add_bullets(s, [
    ("We reproduce Neely et al. (2014) faithfully (99.98% signal match; positive VOL-signal R²_OOS on 1966–2011).", 0),
    ("Extended to 2025, no model — regularized, tree-based, or combined — beats the prevailing mean on RMSE/MAE or direction.", 0),
    ("Economic value persists: switching halves the max drawdown; MV gains of +1.2–1.5% p.a. net of costs.", 0),
    ("Lesson: honest evaluation, disciplined windows and costs matter more than model complexity.", 0),
])
add_notes(s, "Vienna (0:45): wrap-up, hand over to Q&A.")

# ---- 10. Q&A ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
tb = s.shapes.add_textbox(Inches(1), Inches(2.6), Inches(11.3), Inches(1.5))
tf = tb.text_frame
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Thank you — questions?"
r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = DARK
add_notes(s, "All four members (2:00): En Shao on model choices, Sasi on evaluation details, Leon on data, Vienna on interpretation.")

out = os.path.join(ROOT, "deck", "EC4308_Presentation.pptx")
prs.save(out)
print("Wrote", out)
